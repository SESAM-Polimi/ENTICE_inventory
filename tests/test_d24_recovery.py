import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT/'scripts'), str(ROOT/'src')]
from repair_xlsx_directory import recover
from entice_inventory.export.build_d24_inventories import save_workbook_atomic, read_sector_templates
from openpyxl import Workbook, load_workbook


class RecoveryTests(unittest.TestCase):
    def workbook_bytes(self):
        stream = io.BytesIO()
        w = Workbook()
        w.active['A1'] = 'Preserve this value'
        w.active['B1'] = 42
        w.save(stream)
        return stream.getvalue()

    def test_directory_recovery_preserves_every_part_and_cell(self):
        original = self.workbook_bytes()
        directory_start = original.index(b'PK\x01\x02')
        repaired, report = recover(original[:directory_start + 20])
        self.assertEqual(repaired[:directory_start], original[:directory_start])
        with zipfile.ZipFile(io.BytesIO(original)) as a, zipfile.ZipFile(io.BytesIO(repaired)) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for name in a.namelist():
                self.assertEqual(a.read(name), b.read(name))
        w = load_workbook(io.BytesIO(repaired), read_only=True)
        self.assertEqual(w.active['A1'].value, 'Preserve this value')
        self.assertEqual(w.active['B1'].value, 42)
        w.close()

    def test_truncated_payload_is_not_presented_as_recovered(self):
        raw = self.workbook_bytes()
        with self.assertRaises(ValueError):
            recover(raw[:60])

    def test_valid_workbook_is_not_repaired(self):
        with self.assertRaises(ValueError):
            recover(self.workbook_bytes())

    def test_interrupted_save_preserves_existing_destination(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'inventory.xlsx'
            original = self.workbook_bytes()
            path.write_bytes(original)
            w = Workbook()
            def interrupted(temporary):
                Path(temporary).write_bytes(b'incomplete archive')
                raise OSError('Simulated interrupted write')
            with patch.object(w, 'save', side_effect=interrupted):
                with self.assertRaises(OSError):
                    save_workbook_atomic(w, path)
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(Path(folder).iterdir()), [path])

    def test_successful_atomic_save_reopens(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'inventory.xlsx'
            w = Workbook(); w.active['A1'] = 'Complete'
            save_workbook_atomic(w, path)
            check = load_workbook(path, read_only=True)
            self.assertEqual(check.active['A1'].value, 'Complete')
            check.close()

    def test_single_source_selection_ignores_unrelated_broken_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            valid = folder/'Add_sector_Zinc.xlsx'
            w = Workbook(); w.active.title = 'Summary'
            w.active['B2'] = 'ZIR'; w.active['B3'] = 'CHM'
            master = w.create_sheet('Master')
            master.append(['Region', 'Sector']); master.append(['AUT', 'ZIR'])
            w.create_sheet('ZIR_AUT'); w.save(valid)
            (folder/'unrelated.xlsx').write_bytes(b'broken workbook')
            templates, skipped = read_sector_templates(folder, {}, ['ZIR'], [valid])
            self.assertEqual([t.sector_code for t in templates], ['ZIR'])
            self.assertEqual(skipped, [])


if __name__ == '__main__':
    unittest.main()
