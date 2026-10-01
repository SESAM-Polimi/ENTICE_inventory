"""Identity, scope and geography regressions for the canonical registry."""
import copy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'src'))
from entice_inventory.core.registry import Registry, inspect_inventory


class RegistryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.registry = Registry.load()

    def clone(self, sectors=None, regions=None, adapters=None):
        return Registry(sectors or copy.deepcopy(self.registry.sector_document),
                        regions or copy.deepcopy(self.registry.region_document),
                        adapters or copy.deepcopy(list(self.registry.adapters.values())))

    def test_published_coverage_and_inactive_catalogue_entries(self):
        rows = self.registry.sectors
        self.assertEqual(sum(s['lifecycle']=='published_d24' for s in rows.values()),67)
        self.assertEqual(rows['HYS']['lifecycle'],'catalogue_only_not_activated')
        self.assertEqual(rows['PLS']['lifecycle'],'catalogue_only_not_activated')
        self.assertNotEqual(self.registry.sector('NMC')['id'], self.registry.sector('NCM')['id'])
        self.assertNotEqual(self.registry.sector('PLS')['id'], self.registry.sector('PLSR')['id'])

    def test_parent_follows_nace_and_never_legacy_fallback(self):
        for code,parent in [('BVL','MVH'),('MSI','CHM'),('SOP','ELE'),('GAR','NFM'),('NMC','EEQ')]:
            resolved = self.registry.parent(code)
            self.assertEqual(resolved['parent'],parent)
            self.assertTrue(resolved['requires_rebuild'])
        self.assertIsNone(self.registry.parent('APP')['parent'])
        self.assertEqual(self.registry.parent('APP')['published_parent'],'NFM')
        self.assertEqual(self.registry.parent('HYE')['nace_codes'],['20.11'])
        self.assertEqual(self.registry.parent('HYE')['parent'],'CHM')

    def test_scope_is_explicit_and_condition_is_not_hidden(self):
        self.assertIsNone(self.registry.parent('WNT')['parent'])
        self.assertEqual(self.registry.parent('WNT',scope='turbine_sets')['parent'],'OME')
        self.assertEqual(self.registry.parent('WNT',scope='electrical_generators')['parent'],'EEQ')
        self.assertEqual(self.registry.parent('PLR',scope='materials_recovery')['status'],'conditional_on_scope')
        self.assertIsNone(self.registry.parent('PLR')['parent'])
        with self.assertRaises(LookupError):
            self.registry.parent('PLP',scope='not_a_scope')

    def test_aliases_are_namespaced_and_ambiguous_codes_are_rejected(self):
        self.assertEqual(self.registry.sector('PEP')['id'],'PEP')
        # The same text is a source code for PCP in a specific namespace.
        with self.assertRaises(LookupError):
            self.registry.sector('XCH_BPH',namespace='GTAPCE')
        for alias in self.registry.aliases:
            if alias['namespace']=='published_file' and alias['sector_id']=='MUN':
                self.assertEqual(self.registry.sector(alias['value'],'published_file')['id'],'MUN')
        self.assertEqual(self.registry.sector('  bvl ')['id'],'BVL')

    def test_duplicate_identity_or_bridge_never_uses_first_row(self):
        sectors = copy.deepcopy(self.registry.sector_document)
        sectors['sectors'].append(copy.deepcopy(sectors['sectors'][0]))
        with self.assertRaisesRegex(ValueError,'Duplicate'):
            self.clone(sectors=sectors)
        adapters = copy.deepcopy(list(self.registry.adapters.values()))
        adapters[0]['bridge'].append(copy.deepcopy(adapters[0]['bridge'][0]))
        adapters[0]['bridge'][-1]['parent']='WRONG'
        with self.assertRaisesRegex(ValueError,'Duplicate'):
            self.clone(adapters=adapters)

    def test_target_version_and_missing_bridge_do_not_infer_parent(self):
        adapters = copy.deepcopy(list(self.registry.adapters.values()))
        adapters[0]['nace_version']='Rev. 2.1'
        self.assertEqual(self.clone(adapters=adapters).parent('BVL')['status'],'classification_version_mismatch')
        adapters[0]['nace_version']='Rev. 2'
        adapters[0]['bridge']=[r for r in adapters[0]['bridge'] if r['nace_code']!='29.10']
        self.assertEqual(self.clone(adapters=adapters).parent('BVL')['status'],'bridge_missing')
        adapters[0]['bridge'][0]['parent']='TYPO'
        with self.assertRaisesRegex(ValueError,'unknown target parent'):
            self.clone(adapters=adapters)

    def test_global_contains_mauritania_and_cannot_silently_drop_it(self):
        self.assertEqual(len(self.registry.members('GLOBAL')),163)
        self.assertEqual(self.registry.members('Mauritania'),['MRT'])
        self.assertIn('MRT',self.registry.members('GLOBAL'))
        document = copy.deepcopy(self.registry.region_document)
        group = next(g for g in document['geographies'][0]['groups'] if g['id']=='GLOBAL')
        group['members'].remove('MRT')
        with self.assertRaisesRegex(ValueError,'GLOBAL'):
            self.clone(regions=document)

    def test_skipped_checks_are_visible_and_warnings_are_structured(self):
        result = self.registry.check([])
        self.assertEqual(result['executed_checks'],[])
        self.assertEqual(result['skipped_checks'],['classification','identity','regions'])
        for issue in self.registry.check(['classification'])['warnings']:
            self.assertTrue(issue['message'])
            self.assertTrue(issue['suggested_action'])
            self.assertIn('source_record',issue)

    def test_read_only_inventory_check_detects_missing_region_and_wrong_description(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'HYE.xlsx'
            workbook = Workbook()
            master = workbook.active
            master.title='Master'
            master.append(['Region','Sector'])
            master.append(['GLOBAL','HYE'])
            summary = workbook.create_sheet('Summary')
            summary.append(['Sector','HYE','Manufacture of hydrogen via steam reforming'])
            clusters = workbook.create_sheet('Regions Clusters')
            clusters.append(['GLOBAL'])
            for region in self.registry.members('GLOBAL'):
                if region!='MRT':clusters.append([region])
            workbook.save(path)
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            result = inspect_inventory(path,self.registry)
            self.assertEqual(result['coverage']['global_missing_regions'],['MRT'])
            self.assertEqual(result['coverage']['producing_missing_regions'],['MRT'])
            self.assertIn('metadata.sector_description',result['warnings_by_check'])
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)
            only_coverage = inspect_inventory(path,self.registry,['coverage'])
            self.assertEqual(only_coverage['skipped_checks'],['metadata'])

    def test_cli_failure_is_not_success_and_output_is_not_overwritten(self):
        script = Path(__file__).resolve().parents[1]/'scripts/registry.py'
        missing = subprocess.run([sys.executable,str(script),'sector','DOES_NOT_EXIST'],capture_output=True,text=True)
        self.assertEqual(missing.returncode,1)
        self.assertEqual(json.loads(missing.stderr)['status'],'failed')
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)/'check.json'
            output.write_text('preserved')
            result = subprocess.run([sys.executable,str(script),'--output',str(output),'check','--checks'],capture_output=True,text=True)
            self.assertEqual(result.returncode,1)
            self.assertEqual(output.read_text(),'preserved')


if __name__=='__main__':
    unittest.main()
