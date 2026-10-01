"""COPCOSTS pipeline — build Add_sector inventories for the copper sectors.

The secondary-copper sector (CPS) is *not* present in EXIOIOT.csv, so the
standard EXIOIOT pipeline cannot build it. Purdue instead ships the full,
already-normalised cost structure of the three copper sectors in COPCOSTS.csv:

    col 0  cost item   (GTAP-CM v11 sector code, or a factor-of-production code)
    col 1  sector      (cpp / cps / rcp)
    col 2  region      (ita, deu, chn, ...)
    col 3  value       (cost share — sums to 1.0 per sector x region)

This module reads that table for one copper sector, maps every cost item onto a
baseline GTAP12 sector or factor of production, and writes an Add_sector_*.xlsx
workbook in exactly the same format the EXIOIOT pipeline produces, so that
build_d24_inventories.py can consume it unchanged.

For consistency with how CPP/RCP were originally built (EXIOIOT pipeline), only
the sector currently being built resolves to its new ENTICE code; every other
input — including the sibling copper sectors — resolves to its baseline GTAP12
parent (NFM for copper).
"""

from __future__ import annotations

import csv
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl

from entice_inventory.inventory.inventory_writer import write_inventory_workbook
from entice_inventory.core.matching_utils import find_sector_match
from entice_inventory.inventory.exioiot_to_inventory import (
    EPSILON,
    _dedupe,
    _lower,
    _string,
    _upper,
)


def _load_total_outputs(
    splttargs_path: Path,
    repout_sector_code: str,
) -> tuple[list[str], dict[str, float]]:
    """Read REG (region list) and REPOUT (total outputs) from splttargs.xlsx.

    Uses iter_rows rather than random ``ws.cell()`` access, which is
    pathologically slow on a read-only workbook the size of splttargs.
    """
    workbook = openpyxl.load_workbook(
        str(splttargs_path), read_only=True, data_only=True
    )
    if 'REG' not in workbook.sheetnames:
        raise ValueError("splttargs.xlsx is missing the 'REG' sheet.")
    if 'REPOUT' not in workbook.sheetnames:
        raise ValueError("splttargs.xlsx is missing the 'REPOUT' sheet.")

    all_regions: list[str] = []
    for row in workbook['REG'].iter_rows(min_row=2, values_only=True):
        region_code = _upper(row[0]) if row else ''
        if region_code:
            all_regions.append(region_code)

    total_outputs: dict[str, float] = {}
    for row in workbook['REPOUT'].iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        region_code = _upper(row[0])
        sector_code = _lower(row[1]) if len(row) > 1 else ''
        if not region_code or sector_code != repout_sector_code:
            continue
        try:
            value = float(row[2])
        except (TypeError, ValueError, IndexError):
            continue
        total_outputs[region_code] = total_outputs.get(region_code, 0.0) + value

    workbook.close()
    return all_regions, total_outputs


# COPCOSTS cost-item codes that are factors of production, mapped onto the four
# baseline GTAP DB factors (Capital, Labor, TTM, Tax). The matching workbook
# routes every labour occupation to a single "Labor" factor and capital/land/
# natural-resource rents to "Capital".
FACTOR_MAP = {
    'capital': 'Capital',
    'off_mgr_pros': 'Labor',   # Management / officials / professionals
    'tech_aspros': 'Labor',    # Technical and associate professionals
    'clerks': 'Labor',         # Clerical workers
    'service_shop': 'Labor',   # Service and shop workers
    'ag_othlowsk': 'Labor',    # Agriculture and other low-skill workers
}

# Cost-item codes that exist only in the GTAP-CM reference list (no GTAP-ENTICE
# row) and therefore have no automatic GTAP12 mapping. They are all metal flows,
# routed onto the matching baseline GTAP12 sector.
SECTOR_OVERRIDES = {
    'isc': 'I_S',   # Iron and steel casting -> baseline iron & steel
    'rom': 'NFM',   # Recycling - other metals -> non-ferrous metals
    'mps': 'NFM',   # Other metals - secondary -> non-ferrous metals
    'gal': 'NFM',   # Gallium (negligible share) -> non-ferrous metals
}


def _load_resolution_maps(repo_path: Path, exiobase_maps_path: Path):
    """Build the lookup tables used to resolve COPCOSTS cost items.

    Returns (gtap12_baseline, entice2gtap12, gtapce2gtap12, cm2entice) where
    gtap12_baseline maps an upper-cased GTAP12 code to its database casing (most
    sectors are upper-case, but the disaggregated electricity sectors are mixed
    case, e.g. ``CoalBL`` / ``TnD``), and gtapce2gtap12 maps a GTAPCE code to the
    set of GTAP12 codes it covers (a set with more than one member becomes a
    cluster).
    """
    matching_path = repo_path / 'GTAP12_matching.xlsx'
    wb = openpyxl.load_workbook(str(matching_path), data_only=True, read_only=True)
    ws = wb['Sector']

    gtap12_baseline: dict[str, str] = {}
    entice2gtap12: dict[str, str] = {}
    gtapce2gtap12: dict[str, set[str]] = defaultdict(set)
    for row in range(2, ws.max_row + 1):
        entice_code = _lower(ws.cell(row, 2).value)
        gtap12 = _string(ws.cell(row, 3).value)
        gtapce = _lower(ws.cell(row, 4).value)
        if gtap12:
            gtap12_baseline.setdefault(gtap12.upper(), gtap12)
        if entice_code and gtap12:
            entice2gtap12[entice_code] = gtap12
        if gtapce and gtap12:
            gtapce2gtap12[gtapce].add(gtap12)
    wb.close()

    em = openpyxl.load_workbook(
        str(exiobase_maps_path), data_only=True, read_only=True
    )['GTAP_CM_ENTICE']
    cm2entice: dict[str, list[str]] = defaultdict(list)
    for row in em.iter_rows(min_row=2, values_only=True):
        entice_code = _lower(row[0]) if row[0] is not None else ''
        gtap_cm = _lower(row[1]) if len(row) > 1 and row[1] is not None else ''
        if entice_code and gtap_cm:
            cm2entice[gtap_cm].append(entice_code)

    return gtap12_baseline, entice2gtap12, dict(gtapce2gtap12), dict(cm2entice)


def _resolve_input(
    code: str,
    self_code: str,
    new_sector_code: str,
    gtap12_baseline: dict[str, str],
    entice2gtap12: dict[str, str],
    gtapce2gtap12: dict[str, set[str]],
    cm2entice: dict[str, list[str]],
) -> tuple[str, str, list[str] | None]:
    """Resolve a COPCOSTS cost-item code to (mario_type, db_item, cluster).

    cluster is None for a direct mapping, or the list of GTAP12 members when the
    code expands to several baseline sectors.
    """
    key = _lower(code)

    # Factors of production.
    if key in FACTOR_MAP:
        return 'Factor of production', FACTOR_MAP[key], None

    # The sector being built points at its brand-new ENTICE code; sibling copper
    # codes fall through to their baseline GTAP12 parent below.
    if key == self_code:
        return 'Sector', new_sector_code, None

    # Manual overrides for GTAP-CM-only metal flows.
    if key in SECTOR_OVERRIDES:
        return 'Sector', SECTOR_OVERRIDES[key], None

    # Code is already a baseline GTAP12 sector. Return the database casing rather
    # than key.upper(): most sectors are upper-case, but the disaggregated
    # electricity sectors (CoalBL, GasBL, TnD, …) are mixed case and MARIO's
    # sector lookup is case-sensitive.
    if key.upper() in gtap12_baseline:
        return 'Sector', gtap12_baseline[key.upper()], None

    # Code is a GTAP-ENTICE code with a known GTAP12 parent.
    if key in entice2gtap12:
        return 'Sector', entice2gtap12[key], None

    # Code is a GTAPCE code (possibly a cluster over several GTAP12 sectors).
    if key in gtapce2gtap12:
        members = sorted(gtapce2gtap12[key])
        if len(members) == 1:
            return 'Sector', members[0], None
        return 'Sector', _upper(key), members

    # Code is a GTAP-CM code; map through the ENTICE codes it covers.
    if key in cm2entice:
        members = _dedupe(
            entice2gtap12[e] for e in cm2entice[key] if e in entice2gtap12
        )
        if len(members) == 1:
            return 'Sector', members[0], None
        if len(members) > 1:
            return 'Sector', _upper(key), members

    raise ValueError(
        f"No GTAP12 mapping found for COPCOSTS cost item '{code}'. "
        "Add it to FACTOR_MAP, SECTOR_OVERRIDES, or GTAP12_matching.xlsx."
    )


def _validate_copper_sector(copsec_path: Path, sector_key: str) -> None:
    if not copsec_path.exists():
        return
    members: set[str] = set()
    with copsec_path.open(newline='') as handle:
        reader = csv.reader(handle)
        next(reader, None)
        for row in reader:
            if row:
                members.add(_lower(row[0]))
    if members and sector_key not in members:
        raise ValueError(
            f"Sector '{sector_key}' is not listed in COPSEC.csv ({sorted(members)})."
        )


def _load_cost_structure(
    copcosts_path: Path,
    sector_key: str,
    new_sector_code: str,
    maps,
) -> tuple[OrderedDict, OrderedDict, OrderedDict]:
    gtap12_baseline, entice2gtap12, gtapce2gtap12, cm2entice = maps

    inv_by_region_raw: dict[str, dict[tuple[str, str, str], float]] = {}
    used_sector_clusters: OrderedDict[str, list[str]] = OrderedDict()
    used_factor_clusters: OrderedDict[str, list[str]] = OrderedDict()

    with copcosts_path.open(newline='') as handle:
        reader = csv.reader(handle)
        next(reader, None)
        for row in reader:
            if len(row) < 4:
                continue
            cost_item, sector, region, raw_value = row[:4]
            if _lower(sector) != sector_key:
                continue
            try:
                value = float(raw_value)
            except (TypeError, ValueError):
                continue
            if abs(value) <= EPSILON:
                continue

            mario_type, db_item, cluster_members = _resolve_input(
                cost_item,
                sector_key,
                new_sector_code,
                gtap12_baseline,
                entice2gtap12,
                gtapce2gtap12,
                cm2entice,
            )
            if cluster_members:
                if mario_type == 'Sector':
                    used_sector_clusters.setdefault(db_item, cluster_members)
                else:
                    used_factor_clusters.setdefault(db_item, cluster_members)

            region_code = _upper(region)
            agg = inv_by_region_raw.setdefault(region_code, {})
            key = (mario_type, db_item, 'GLOBAL')
            agg[key] = agg.get(key, 0.0) + value

    if not inv_by_region_raw:
        raise ValueError(
            f"No COPCOSTS rows found for sector '{sector_key.upper()}'."
        )

    inv_by_region: OrderedDict[str, OrderedDict] = OrderedDict()
    for region_code in sorted(inv_by_region_raw.keys()):
        agg = inv_by_region_raw[region_code]
        total = sum(agg.values())
        if abs(total) <= EPSILON:
            raise ValueError(
                f"The cost structure for region '{region_code}' has zero total."
            )
        normalized = OrderedDict((k, v / total) for k, v in agg.items())
        inv_by_region[region_code] = normalized

    return inv_by_region, used_sector_clusters, used_factor_clusters


def make_inventory(
    inventory_name: str,
    repo_path: str | Path,
    copcosts_path: str | Path,
    copsec_path: str | Path,
    splttargs_path: str | Path,
    exiobase_maps_path: str | Path,
    output_path: str | Path = None,
    version: str = '',
    residual_other_sectors_share: float = 0.0,
):
    repo_path = Path(repo_path)
    copcosts_path = Path(copcosts_path)
    copsec_path = Path(copsec_path)
    splttargs_path = Path(splttargs_path)
    exiobase_maps_path = Path(exiobase_maps_path)

    target = find_sector_match(repo_path, inventory_name)
    if target is None:
        raise ValueError(
            f"Sector '{inventory_name}' was not found in GTAP12_matching.xlsx."
        )
    if _lower(target.pipeline) != 'copcosts':
        raise ValueError(
            f"Sector '{inventory_name}' is not marked with Pipeline = COPCOSTS."
        )
    if not target.entice_code or not target.gtap12:
        raise ValueError(
            f"Sector '{inventory_name}' is missing GTAP ENTICE code or GTAP12."
        )

    sector_code = _string(target.entice_code)
    parent_code = _string(target.gtap12)
    sector_key = _lower(sector_code)
    target_repout_code = _lower(sector_code)

    output_path = (
        Path(output_path)
        if output_path is not None
        else repo_path / f"Add_sector_{inventory_name}.xlsx"
    )

    _validate_copper_sector(copsec_path, sector_key)

    maps = _load_resolution_maps(repo_path, exiobase_maps_path)
    inv_by_region, used_sector_clusters, used_factor_clusters = _load_cost_structure(
        copcosts_path, sector_key, sector_code, maps
    )

    all_regions, total_outputs = _load_total_outputs(splttargs_path, target_repout_code)
    if not all_regions:
        raise ValueError("No GTAP regions found in splttargs.xlsx/REG.")

    missing_cost_regions = [r for r in all_regions if r not in inv_by_region]
    if missing_cost_regions:
        print(
            "  Warning: COPCOSTS has no cost structure for "
            f"{len(missing_cost_regions)} region(s): {missing_cost_regions[:10]}"
        )

    inventory_regions = [r for r in all_regions if r in inv_by_region]
    inventory_regions.extend(r for r in inv_by_region if r not in set(all_regions))

    print(f"  New sector : {inventory_name}  [{sector_code}]")
    print(f"  Parent     : {parent_code}")
    print(f"  COPCOSTS key    : {sector_key.upper()}")
    print(f"  REPOUT output   : {target_repout_code.upper()}")
    print(f"  Cost structures : {len(inventory_regions)} region(s)")
    print(f"  Total outputs   : {len(total_outputs)} non-zero region(s) from REPOUT")
    print(f"  Sector clusters : {list(used_sector_clusters.keys())}")
    print(f"  Factor clusters : {list(used_factor_clusters.keys())}")

    sources = [
        ('COPCOSTS', str(copcosts_path)),
        ('COPSEC', str(copsec_path)),
        ('splttargs / REPOUT', str(splttargs_path)),
    ]

    return write_inventory_workbook(
        output_path,
        inventory_name=inventory_name,
        sector_code=sector_code,
        parent_code=parent_code,
        version=version,
        data_collection_lead='COPCOSTS pipeline',
        sources=sources,
        all_regions=all_regions,
        inventory_regions=inventory_regions,
        inv_by_region=inv_by_region,
        total_outputs=total_outputs,
        used_sector_clusters=used_sector_clusters,
        used_factor_clusters=used_factor_clusters,
        repo_path=repo_path,
        residual_other_sectors_share=residual_other_sectors_share,
    )
