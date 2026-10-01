"""EXIOIOT pipeline — build Add_sector inventories from EXIOIOT.csv.

For sectors flagged ``Pipeline = EXIOIOT`` in GTAP12_matching.xlsx, this module
reads the already-normalised EXIOIOT cost structure (one row per
cost-item x region) for the target GTAPCE code, maps each input onto a baseline
GTAP12 sector/cluster or factor of production, pulls regional outputs from
splttargs.xlsx (REPOUT), and writes an ``Add_sector_*.xlsx`` workbook via
``inventory_writer.write_inventory_workbook`` so that build_d24_inventories.py
can consume it unchanged.

API
---
  from exioiot_to_inventory import make_inventory
  make_inventory("Manufacture of primary copper", repo_path, exioiot_path,
                 splttargs_path, version="Y26M05")
"""

from __future__ import annotations

import csv
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl

from entice_inventory.inventory.inventory_writer import write_inventory_workbook
from entice_inventory.core.matching_utils import find_sector_match, load_matching_rows


VALUE_ADDED_CODES = {
    'fixcap',
    'labhskl',
    'lablskl',
    'labmskl',
    'othrent',
    'othtax',
    'prodtax',
}

FACTOR_FALLBACKS = {
    # EXIOIOT aggregates rents into OTHRENT, while the matching workbook
    # splits land and natural resources into GTAPCE categories not present here.
    'othrent': ['Capital'],
}

EPSILON = 1e-12


def _string(value) -> str:
    return str(value).strip() if value is not None else ''


def _lower(value) -> str:
    return _string(value).lower()


def _upper(value) -> str:
    return _string(value).upper()


def _dedupe(values) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for value in values:
        if value and value not in seen:
            ordered.append(value)
            seen.add(value)
    return ordered


def _build_sector_resolution(sector_rows) -> tuple[dict[str, str], dict[str, list[str]]]:
    direct_from_entice_code: dict[str, str] = {}
    grouped_gtap12: dict[str, list[str]] = defaultdict(list)

    for row in sector_rows:
        gtapce = _lower(row.gtapce)
        entice_code = _lower(row.entice_code)
        if not gtapce:
            gtap12_code = _string(row.gtap12)
            if entice_code and gtap12_code:
                direct_from_entice_code[entice_code] = gtap12_code
            continue

        gtap12_code = _string(row.gtap12)
        if gtap12_code:
            grouped_gtap12[gtapce].append(gtap12_code)

        if entice_code and gtap12_code:
            direct_from_entice_code[entice_code] = gtap12_code

    direct_map: dict[str, str] = dict(direct_from_entice_code)
    cluster_map: dict[str, list[str]] = {}
    for gtapce, gtap12_codes in grouped_gtap12.items():
        members = _dedupe(gtap12_codes)
        if len(members) == 1:
            direct_map[gtapce] = members[0]
        elif len(members) > 1:
            cluster_map[gtapce] = members

    return direct_map, cluster_map


def _build_factor_resolution(factor_rows) -> tuple[dict[str, str], dict[str, list[str]]]:
    grouped_gtap12: dict[str, list[str]] = defaultdict(list)
    for row in factor_rows:
        gtapce = _lower(row.gtapce)
        gtap12_code = _string(row.gtap12)
        if gtapce and gtap12_code:
            grouped_gtap12[gtapce].append(gtap12_code)

    for gtapce, members in FACTOR_FALLBACKS.items():
        grouped_gtap12.setdefault(gtapce, []).extend(members)

    direct_map: dict[str, str] = {}
    cluster_map: dict[str, list[str]] = {}
    for gtapce, gtap12_codes in grouped_gtap12.items():
        members = _dedupe(gtap12_codes)
        if len(members) == 1:
            direct_map[gtapce] = members[0]
        elif len(members) > 1:
            cluster_map[gtapce] = members

    return direct_map, cluster_map

def _resolve_sector_input(
    gtapce_code: str,
    direct_map: dict[str, str],
    cluster_map: dict[str, list[str]],
) -> tuple[str, list[str] | None]:
    code = _lower(gtapce_code)
    if code in direct_map:
        return direct_map[code], None
    if code in cluster_map:
        return _upper(code), cluster_map[code]
    raise ValueError(f"No sector mapping found for GTAPCE code '{gtapce_code}'.")


def _resolve_factor_input(
    gtapce_code: str,
    direct_map: dict[str, str],
    cluster_map: dict[str, list[str]],
) -> tuple[str, list[str] | None]:
    code = _lower(gtapce_code)
    if code in direct_map:
        return direct_map[code], None
    if code in cluster_map:
        return _upper(code), cluster_map[code]
    raise ValueError(f"No factor mapping found for GTAPCE code '{gtapce_code}'.")


def _collect_target_sector_codes(exioiot_path: Path, target_gtapce: str) -> set[str]:
    sector_codes: set[str] = set()
    with exioiot_path.open(newline='') as handle:
        reader = csv.reader(handle)
        next(reader, None)
        for row in reader:
            if len(row) < 4:
                continue
            gcerows, mutsec = row[0], row[1]
            if _lower(mutsec) != target_gtapce:
                continue
            gcerows_code = _lower(gcerows)
            if gcerows_code and gcerows_code not in VALUE_ADDED_CODES:
                sector_codes.add(gcerows_code)
    return sector_codes


def _unresolved_sector_codes(
    raw_codes: set[str],
    direct_map: dict[str, str],
    cluster_map: dict[str, list[str]],
) -> list[str]:
    return sorted(
        code for code in raw_codes
        if code not in direct_map and code not in cluster_map
    )


def _load_target_cost_structure(
    exioiot_path: Path,
    target_gtapce: str,
    sector_direct_map: dict[str, str],
    sector_cluster_map: dict[str, list[str]],
    factor_direct_map: dict[str, str],
    factor_cluster_map: dict[str, list[str]],
) -> tuple[OrderedDict, OrderedDict, OrderedDict]:
    inv_by_region_raw: dict[str, dict[tuple[str, str, str], float]] = {}
    used_sector_clusters: OrderedDict[str, list[str]] = OrderedDict()
    used_factor_clusters: OrderedDict[str, list[str]] = OrderedDict()

    with exioiot_path.open(newline='') as handle:
        reader = csv.reader(handle)
        next(reader, None)
        for row in reader:
            if len(row) < 4:
                continue

            gcerows, mutsec, region, raw_value = row[:4]
            if _lower(mutsec) != target_gtapce:
                continue

            try:
                value = float(raw_value)
            except (TypeError, ValueError):
                continue

            if abs(value) <= EPSILON:
                continue

            region_code = _upper(region)
            gcerows_code = _lower(gcerows)

            if gcerows_code in VALUE_ADDED_CODES:
                mario_type = 'Factor of production'
                db_item, cluster_members = _resolve_factor_input(
                    gcerows_code,
                    factor_direct_map,
                    factor_cluster_map,
                )
                if cluster_members:
                    used_factor_clusters.setdefault(db_item, cluster_members)
            else:
                mario_type = 'Sector'
                db_item, cluster_members = _resolve_sector_input(
                    gcerows_code,
                    sector_direct_map,
                    sector_cluster_map,
                )
                if cluster_members:
                    used_sector_clusters.setdefault(db_item, cluster_members)

            agg = inv_by_region_raw.setdefault(region_code, {})
            key = (mario_type, db_item, 'GLOBAL')
            agg[key] = agg.get(key, 0.0) + value

    if not inv_by_region_raw:
        raise ValueError(
            f"No EXIOIOT rows found for target GTAPCE code '{target_gtapce.upper()}'."
        )

    inv_by_region: OrderedDict[str, OrderedDict] = OrderedDict()
    for region_code in sorted(inv_by_region_raw.keys()):
        agg = inv_by_region_raw[region_code]
        total = sum(agg.values())
        if abs(total) <= EPSILON:
            raise ValueError(
                f"The normalized cost structure for region '{region_code}' has zero total."
            )

        normalized = OrderedDict()
        for key, value in agg.items():
            normalized[key] = value / total
        inv_by_region[region_code] = normalized

    return inv_by_region, used_sector_clusters, used_factor_clusters


def _load_total_outputs(
    splttargs_path: Path,
    repout_sector_code: str,
) -> tuple[list[str], dict[str, float]]:
    workbook = openpyxl.load_workbook(str(splttargs_path), read_only=True, data_only=True)

    if 'REG' not in workbook.sheetnames:
        raise ValueError("splttargs.xlsx is missing the 'REG' sheet.")
    if 'REPOUT' not in workbook.sheetnames:
        raise ValueError("splttargs.xlsx is missing the 'REPOUT' sheet.")

    all_regions = []
    ws_reg = workbook['REG']
    for row in range(2, ws_reg.max_row + 1):
        region_code = _upper(ws_reg.cell(row, 1).value)
        if region_code:
            all_regions.append(region_code)

    total_outputs: dict[str, float] = {}
    ws_repout = workbook['REPOUT']
    for row in range(2, ws_repout.max_row + 1):
        region_code = _upper(ws_repout.cell(row, 1).value)
        sector_code = _lower(ws_repout.cell(row, 2).value)
        raw_value = ws_repout.cell(row, 3).value
        if not region_code or sector_code != repout_sector_code:
            continue
        try:
            value = float(raw_value)
        except (TypeError, ValueError):
            continue
        total_outputs[region_code] = total_outputs.get(region_code, 0.0) + value

    return all_regions, total_outputs


def make_inventory(
    inventory_name: str,
    repo_path: str | Path,
    exioiot_path: str | Path,
    splttargs_path: str | Path,
    output_path: str | Path = None,
    version: str = '',
    residual_other_sectors_share: float = 0.0,
):
    repo_path = Path(repo_path)
    exioiot_path = Path(exioiot_path)
    splttargs_path = Path(splttargs_path)

    target = find_sector_match(repo_path, inventory_name)
    if target is None:
        raise ValueError(
            f"Sector '{inventory_name}' was not found in GTAP12_matching.xlsx."
        )
    if _lower(target.pipeline) != 'exioiot':
        raise ValueError(
            f"Sector '{inventory_name}' is not marked with Pipeline = EXIOIOT."
        )
    if not target.entice_code or not target.gtapce or not target.gtap12:
        raise ValueError(
            f"Sector '{inventory_name}' is missing GTAP ENTICE code, GTAPCE, or GTAP12."
        )

    sector_code = _string(target.entice_code)
    parent_code = _string(target.gtap12)
    target_gtapce = _lower(target.gtapce)
    target_repout_code = _lower(sector_code)

    output_path = (
        Path(output_path)
        if output_path is not None
        else repo_path / f"Add_sector_{inventory_name}.xlsx"
    )

    sector_rows, factor_rows = load_matching_rows(repo_path)
    sector_direct_map, sector_cluster_map = _build_sector_resolution(sector_rows)
    factor_direct_map, factor_cluster_map = _build_factor_resolution(factor_rows)

    # Only the sector currently being added should resolve to the new ENTICE code.
    # Other EXIOIOT sectors that appear as inputs must resolve to existing GTAP12
    # sectors/clusters, otherwise MARIO cannot find them in the baseline database.
    sector_direct_map[target_gtapce] = sector_code
    sector_direct_map[_lower(sector_code)] = sector_code
    sector_cluster_map.pop(target_gtapce, None)
    sector_cluster_map.pop(_lower(sector_code), None)

    raw_sector_codes = _collect_target_sector_codes(exioiot_path, target_gtapce)
    unresolved_codes = _unresolved_sector_codes(
        raw_sector_codes,
        sector_direct_map,
        sector_cluster_map,
    )
    if unresolved_codes:
        raise ValueError(
            "Missing sector mappings for EXIOIOT inputs in GTAP12_matching: "
            + ', '.join(unresolved_codes[:20])
        )

    inv_by_region, used_sector_clusters, used_factor_clusters = _load_target_cost_structure(
        exioiot_path,
        target_gtapce,
        sector_direct_map,
        sector_cluster_map,
        factor_direct_map,
        factor_cluster_map,
    )

    all_regions, total_outputs = _load_total_outputs(splttargs_path, target_repout_code)
    if not all_regions:
        raise ValueError("No GTAP regions found in splttargs.xlsx/REG.")

    missing_cost_regions = [region for region in all_regions if region not in inv_by_region]
    if missing_cost_regions:
        print(
            "  Warning: EXIOIOT has no cost structure for "
            f"{len(missing_cost_regions)} region(s): {missing_cost_regions[:10]}"
        )

    inventory_regions = [region for region in all_regions if region in inv_by_region]
    extra_regions = [region for region in inv_by_region if region not in set(all_regions)]
    inventory_regions.extend(extra_regions)

    print(f"  New sector : {inventory_name}  [{sector_code}]")
    print(f"  Parent     : {parent_code}")
    print(f"  GTAPCE cost     : {target_gtapce.upper()}")
    print(f"  REPOUT output   : {target_repout_code.upper()}")
    print(f"  Cost structures : {len(inventory_regions)} region(s)")
    print(f"  Total outputs   : {len(total_outputs)} non-zero region(s) from REPOUT")
    print(f"  Sector clusters : {list(used_sector_clusters.keys())}")
    print(f"  Factor clusters : {list(used_factor_clusters.keys())}")

    sources = [
        ('EXIOIOT', str(exioiot_path)),
        ('splttargs / REPOUT', str(splttargs_path)),
    ]

    return write_inventory_workbook(
        output_path,
        inventory_name=inventory_name,
        sector_code=sector_code,
        parent_code=parent_code,
        version=version,
        data_collection_lead='EXIOIOT pipeline',
        sources=sources,
        all_regions=all_regions,
        inventory_regions=inventory_regions,
        inv_by_region=inv_by_region,
        total_outputs=total_outputs,
        used_sector_clusters=used_sector_clusters,
        used_factor_clusters=used_factor_clusters,
        repo_path=repo_path,
        residual_other_sectors_share=residual_other_sectors_share,
        print_residual_detail=True,
    )
