"""Shared writer for Add_sector inventory workbooks.

The EXIOIOT and COPCOSTS pipelines resolve their cost structures very
differently, but both emit the *same* ``Add_sector_*.xlsx`` layout that
``build_d24_inventories.py`` consumes. This module owns that layout in one
place so the two pipelines stay byte-for-byte consistent.

Sheets written (in order):
    Summary            sector metadata, sources, page links
    Master             one row per producing region (Split operation)
    Regions Clusters   GLOBAL = all GTAP regions
    Sectors Clusters   input-sector clusters actually used
    Factors Clusters   factor clusters actually used
    {sector_code}_{region}   one regional inventory sheet per region
    DB units           static unit table
    Total outputs      monetary output per region

Residual "rest of parent" clusters are appended last via
``residual_sector_cluster.apply_residual_sector_clusters``.
"""

from __future__ import annotations

from collections import OrderedDict
from pathlib import Path

import openpyxl

from entice_inventory.core.residual_sector_cluster import apply_residual_sector_clusters
from entice_inventory.inventory.template_to_inventory import DB_UNITS_DATA, _safe_sheet_name, _write_clusters_sheet


def write_inventory_workbook(
    output_path: str | Path,
    *,
    inventory_name: str,
    sector_code: str,
    parent_code: str,
    version: str,
    data_collection_lead: str,
    sources: list[tuple[str, str]],
    all_regions: list[str],
    inventory_regions: list[str],
    inv_by_region: dict[str, "OrderedDict[tuple[str, str, str], float]"],
    total_outputs: dict[str, float],
    used_sector_clusters: "OrderedDict[str, list[str]]",
    used_factor_clusters: "OrderedDict[str, list[str]]",
    repo_path: str | Path,
    residual_other_sectors_share: float = 0.0,
    print_residual_detail: bool = False,
) -> Path:
    """Write one Add_sector inventory workbook and return its path.

    ``data_collection_lead`` and ``sources`` identify the producing pipeline;
    every other sheet is identical regardless of the source pipeline.
    """
    output_path = Path(output_path)
    repo_path = Path(repo_path)

    regions_clusters = OrderedDict([('GLOBAL', all_regions)])

    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)
    all_sheet_names: list[str] = []

    ws_sum = wb_out.create_sheet('Summary')
    ws_sum.cell(1, 1).value = 'INFO'
    ws_sum.cell(2, 1).value = 'Sector'
    ws_sum.cell(2, 2).value = sector_code
    ws_sum.cell(2, 3).value = inventory_name
    ws_sum.cell(3, 1).value = 'Parent sector'
    ws_sum.cell(3, 2).value = parent_code
    ws_sum.cell(4, 1).value = 'Inventory version'
    ws_sum.cell(4, 2).value = version
    ws_sum.cell(5, 1).value = 'Data collection lead'
    ws_sum.cell(5, 2).value = data_collection_lead
    ws_sum.cell(7, 1).value = 'Sources'
    for idx, (source_name, source_url) in enumerate(sources, start=8):
        ws_sum.cell(idx, 1).value = source_name
        ws_sum.cell(idx, 2).value = source_url

    ws_master = wb_out.create_sheet('Master')
    all_sheet_names.append('Master')
    ws_master.append([
        'Region', 'Sector', 'Inventory sheet', 'Quantity', 'Unit',
        'Final consumption', 'Consumption category',
        'Parent Sector', 'Leave empty', 'Source', 'Notes', 'Add or Split',
    ])
    for region_code in inventory_regions:
        sheet_name = _safe_sheet_name(f"{sector_code}_{region_code}")
        ws_master.append([
            region_code, sector_code, sheet_name, None, 'M USD',
            None, None, parent_code, None, None, None, 'Split',
        ])

    _write_clusters_sheet(wb_out.create_sheet('Regions Clusters'), regions_clusters)
    all_sheet_names.append('Regions Clusters')

    _write_clusters_sheet(wb_out.create_sheet('Sectors Clusters'), used_sector_clusters)
    all_sheet_names.append('Sectors Clusters')

    _write_clusters_sheet(wb_out.create_sheet('Factors Clusters'), used_factor_clusters)
    all_sheet_names.append('Factors Clusters')

    for region_code in inventory_regions:
        sheet_name = _safe_sheet_name(f"{sector_code}_{region_code}")
        ws_inv = wb_out.create_sheet(sheet_name)
        ws_inv.append([
            'Quantity', 'Unit', 'Input', 'Item type',
            'DB Item', 'DB Region', 'Change type',
        ])
        for (mario_type, db_item, db_region), value in inv_by_region[region_code].items():
            ws_inv.append([value, 'M USD', None, mario_type, db_item, db_region, 'Update'])
        all_sheet_names.append(sheet_name)

    ws_db = wb_out.create_sheet('DB units')
    for row in DB_UNITS_DATA:
        ws_db.append(list(row))
    all_sheet_names.append('DB units')

    ws_to = wb_out.create_sheet('Total outputs')
    ws_to.append(['Sector', 'Region', 'Quantity', 'Unit', 'Source', 'Notes'])
    for region_code in all_regions:
        ws_to.append([sector_code, region_code, total_outputs.get(region_code, 0.0), 'M USD'])
    all_sheet_names.append('Total outputs')

    ws_sum.cell(1, 9).value = 'LINKS TO PAGES'
    for idx, sheet_name in enumerate(all_sheet_names, start=2):
        cell = ws_sum.cell(idx, 9)
        cell.value = sheet_name
        safe_name = sheet_name.replace("'", "''")
        cell.hyperlink = f"#'{safe_name}'!A1"

    residual_updates = apply_residual_sector_clusters(
        wb_out,
        repo_path,
        residual_other_sectors_share,
        sector_code=sector_code,
    )
    if residual_updates:
        print(
            "  Residual sector clusters added : "
            f"{len(residual_updates)} inventory sheet(s) "
            f"at {residual_other_sectors_share:.2%}"
        )
        if print_residual_detail:
            for update in residual_updates[:10]:
                print(
                    f"    {update.sheet_name}: {update.cluster_name} "
                    f"({len(update.cluster_members)} member(s), "
                    f"{len(update.missing_codes)} missing code(s))"
                )

    wb_out.save(str(output_path))
    print(f"\nInventory saved: {output_path.name}")
    return output_path
