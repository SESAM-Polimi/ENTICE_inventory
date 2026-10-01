from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

from entice_inventory.core.paths import find_data_file


@dataclass(frozen=True)
class SectorMatch:
    entice_name: str
    entice_code: str
    gtap12: str
    gtapce: str
    pipeline: str


@dataclass(frozen=True)
class FactorMatch:
    factor_name: str
    gtap12: str
    gtapce: str


def _string(value) -> str:
    return str(value).strip() if value is not None else ''


def _header_key(value) -> str:
    return _string(value).casefold()


def _resolve_matching_path(path_or_repo: str | Path) -> Path:
    path = Path(path_or_repo)
    if path.is_dir():
        return find_data_file('GTAP12_matching.xlsx', base=path)
    return path


def _header_indices(ws) -> dict[str, int]:
    return {
        _header_key(ws.cell(1, column).value): column
        for column in range(1, ws.max_column + 1)
        if _string(ws.cell(1, column).value)
    }


def _require_header(ws, header_name: str) -> int:
    header_map = _header_indices(ws)
    column = header_map.get(header_name.casefold())
    if column is None:
        raise ValueError(
            f"Sheet '{ws.title}' in GTAP12_matching.xlsx is missing the "
            f"'{header_name}' column."
        )
    return column


def _load_firstcol_to_namedcol(ws, named_header: str) -> dict[str, str]:
    value_column = _require_header(ws, named_header)
    mapping: dict[str, str] = {}
    for row in range(2, ws.max_row + 1):
        key = _string(ws.cell(row, 1).value)
        value = _string(ws.cell(row, value_column).value)
        if key and value:
            mapping[key] = value
    return mapping


def load_legacy_gtap12_maps(path_or_repo: str | Path) -> tuple[dict[str, str], dict[str, str]]:
    matching_path = _resolve_matching_path(path_or_repo)
    workbook = openpyxl.load_workbook(str(matching_path), data_only=True)

    if 'Sector' not in workbook.sheetnames:
        raise ValueError("GTAP12_matching.xlsx is missing the 'Sector' sheet.")
    if 'Factor of production' not in workbook.sheetnames:
        raise ValueError(
            "GTAP12_matching.xlsx is missing the 'Factor of production' sheet."
        )

    sector_map = _load_firstcol_to_namedcol(workbook['Sector'], 'GTAP12')
    factor_map = _load_firstcol_to_namedcol(workbook['Factor of production'], 'GTAP12')
    return sector_map, factor_map


def load_matching_rows(path_or_repo: str | Path) -> tuple[list[SectorMatch], list[FactorMatch]]:
    matching_path = _resolve_matching_path(path_or_repo)
    workbook = openpyxl.load_workbook(str(matching_path), data_only=True)

    if 'Sector' not in workbook.sheetnames:
        raise ValueError("GTAP12_matching.xlsx is missing the 'Sector' sheet.")
    if 'Factor of production' not in workbook.sheetnames:
        raise ValueError(
            "GTAP12_matching.xlsx is missing the 'Factor of production' sheet."
        )

    sector_ws = workbook['Sector']
    sector_headers = _header_indices(sector_ws)
    sector_code_col = sector_headers.get('gtap entice codes')
    sector_gtap12_col = _require_header(sector_ws, 'GTAP12')
    sector_gtapce_col = sector_headers.get('gtapce')
    sector_pipeline_col = sector_headers.get('pipeline')

    sector_rows: list[SectorMatch] = []
    for row in range(2, sector_ws.max_row + 1):
        entice_name = _string(sector_ws.cell(row, 1).value)
        if not entice_name:
            continue
        sector_rows.append(
            SectorMatch(
                entice_name=entice_name,
                entice_code=_string(sector_ws.cell(row, sector_code_col).value)
                if sector_code_col
                else '',
                gtap12=_string(sector_ws.cell(row, sector_gtap12_col).value),
                gtapce=_string(sector_ws.cell(row, sector_gtapce_col).value)
                if sector_gtapce_col
                else '',
                pipeline=_string(sector_ws.cell(row, sector_pipeline_col).value)
                if sector_pipeline_col
                else '',
            )
        )

    factor_ws = workbook['Factor of production']
    factor_gtap12_col = _require_header(factor_ws, 'GTAP12')
    factor_headers = _header_indices(factor_ws)
    factor_gtapce_col = factor_headers.get('gtapce')

    factor_rows: list[FactorMatch] = []
    for row in range(2, factor_ws.max_row + 1):
        factor_name = _string(factor_ws.cell(row, 1).value)
        if not factor_name:
            continue
        factor_rows.append(
            FactorMatch(
                factor_name=factor_name,
                gtap12=_string(factor_ws.cell(row, factor_gtap12_col).value),
                gtapce=_string(factor_ws.cell(row, factor_gtapce_col).value)
                if factor_gtapce_col
                else '',
            )
        )

    return sector_rows, factor_rows


def load_parent_map(path_or_repo: str | Path) -> dict[str, str]:
    """Legacy sub-sector code -> parent GTAP12 code, for historical pipelines.

    Built from the 'Sector' sheet of GTAP12_matching.xlsx
    ('GTAP ENTICE CODES' -> 'GTAP12'). This workbook is not NACE certification.
    New consumers should use Registry.parent(), which exposes unresolved scopes
    and never falls back to these labels. Codes are upper-cased; historical
    callers retain the first occurrence of each code.
    """
    sector_rows, _ = load_matching_rows(path_or_repo)
    parents: dict[str, str] = {}
    for row in sector_rows:
        code = row.entice_code.strip().upper()
        parent = row.gtap12.strip()
        if code and parent:
            parents.setdefault(code, parent)
    return parents


def find_sector_match(
    path_or_repo: str | Path,
    inventory_name: str,
) -> SectorMatch | None:
    needle = inventory_name.strip().casefold()
    sector_rows, _ = load_matching_rows(path_or_repo)
    for row in sector_rows:
        if row.entice_name.casefold() == needle:
            return row
    return None


def find_sector_matches_by_code(
    path_or_repo: str | Path,
    sector_code: str,
) -> list[SectorMatch]:
    needle = sector_code.strip().upper()
    if not needle:
        return []

    sector_rows, _ = load_matching_rows(path_or_repo)
    return [
        row for row in sector_rows
        if row.entice_code.strip().upper() == needle
    ]


def find_sector_match_by_code(
    path_or_repo: str | Path,
    sector_code: str,
) -> SectorMatch:
    matches = find_sector_matches_by_code(path_or_repo, sector_code)
    if not matches:
        raise ValueError(
            f"No row in GTAP12_matching.xlsx has GTAP ENTICE CODES = '{sector_code}'."
        )
    if len(matches) > 1:
        names = '; '.join(row.entice_name for row in matches)
        raise ValueError(
            f"GTAP ENTICE code '{sector_code}' is ambiguous in GTAP12_matching.xlsx: {names}"
        )
    return matches[0]
