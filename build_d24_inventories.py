from __future__ import annotations

import argparse
import importlib
import math
import sys
import warnings
from collections import OrderedDict, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


REPO_DIR = Path(__file__).resolve().parent
DEFAULT_DATA_ROOT = (
    Path.home()
    / "Library/CloudStorage/OneDrive-SharedLibraries-eNextGen/ENTICE - Documents/WPs, Tasks & Deliverables/WP2 - Data/T2.2 & T2.3 - GTAP disaggregation"
)
DEFAULT_DB_PATH = DEFAULT_DATA_ROOT / "Database/GTAP 2023/2023entice"
DEFAULT_SOURCE_DIR = DEFAULT_DATA_ROOT / "Data collection/Inventory cleaning/MARIO inventories copy"
DEFAULT_OUTPUT_DIR = DEFAULT_DATA_ROOT / "Data collection/Inventory cleaning/D2.4 inventories"
DEFAULT_PURDUE_SPLITARGS_PATH = DEFAULT_DATA_ROOT / "Shared material/Purdue data collection/June1/splttargs.xlsx"
DEFAULT_GTAP_SECTORS_PATH = DEFAULT_DATA_ROOT / "Shared material/Purdue data collection/May13/GTAP sectors H5.xlsx"
DEFAULT_MARIO_SRC = Path.home() / "Documents/GitHub/MARIO"
REPORT_FILENAME = "export_d24_report.txt"
DEFAULT_INVENTORY_SUM_CHECK_TOLERANCE = 1e-3
STATIC_SHEETS = {
    "Summary",
    "Master",
    "Regions Clusters",
    "Sectors Clusters",
    "Factors Clusters",
    "Trades",
    "DB units",
    "Total outputs",
}


@dataclass
class SectorTemplate:
    source_path: Path
    sector_code: str
    sector_name: str
    parent_sector: str
    inventory_version: str | None
    data_collection_lead: str | None
    add_or_split: str
    summary_sources: list[tuple[str | None, str | None]]
    parent_sector_name: str | None = None


def as_path(value: str | Path) -> Path:
    return value if isinstance(value, Path) else Path(value)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Build D2.4-ready sector inventory workbooks by parsing GTAP, "
            "adding the MARIO inventory sectors, and exporting "
            "regional cost structures."
        )
    )
    parser.add_argument("--db-path", type=Path, default=DEFAULT_DB_PATH)
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--purdue-splitargs-path", type=Path, default=DEFAULT_PURDUE_SPLITARGS_PATH)
    parser.add_argument("--mario-src", type=Path, default=DEFAULT_MARIO_SRC)
    parser.add_argument(
        "--tolerance",
        type=float,
        default=1e-12,
        help="Absolute threshold below which coefficients are omitted from inventory sheets.",
    )
    parser.add_argument(
        "--inventory-sum-check-tolerance",
        type=float,
        default=DEFAULT_INVENTORY_SUM_CHECK_TOLERANCE,
        help=(
            "Absolute tolerance for flagging exported inventories whose written rows do not sum to 1. "
            "Default: 1e-3, equivalent to accepting totals in the [0.999, 1.001] range."
        ),
    )
    return parser.parse_args()


def ensure_mario_import(mario_src: str | Path):
    mario_src = as_path(mario_src)
    candidates: list[Path] = []
    if mario_src:
        candidates.append(mario_src)
    fallback = DEFAULT_MARIO_SRC
    if fallback not in candidates:
        candidates.append(fallback)

    for candidate in candidates:
        if candidate.exists() and str(candidate) not in sys.path:
            sys.path.insert(0, str(candidate))

    mario = importlib.import_module("mario")
    if hasattr(mario, "parse_from_parquet"):
        return mario

    for candidate in candidates:
        if candidate.exists() and str(candidate) not in sys.path:
            sys.path.insert(0, str(candidate))
            importlib.invalidate_caches()
            mario = importlib.reload(mario)
            if hasattr(mario, "parse_from_parquet"):
                return mario

    raise RuntimeError(
        "The available 'mario' module does not expose parse_from_parquet. "
        f"Checked fallback source path '{mario_src}'."
    )


def workbook_paths(source_dir: str | Path) -> list[Path]:
    source_dir = as_path(source_dir)
    return sorted(path for path in source_dir.glob("*.xlsx") if not path.name.startswith("~$"))


def sector_name_from_source_path(source_path: str | Path) -> str:
    stem = as_path(source_path).stem
    if stem.startswith("Add_sector_"):
        return stem.removeprefix("Add_sector_")
    return stem


def normalize_summary_sources(
    summary_sources: list[tuple[str | None, str | None]],
) -> list[tuple[str | None, str | None]]:
    labels = {
        str(label).strip().upper()
        for label, _ in summary_sources
        if label is not None
    }
    if {"EXIOIOT", "SPLTTARGS / REPOUT"}.issubset(labels):
        return [("EXIOBASE", None), ("GTAP CE", None), ("GTAP CM", None)]
    return summary_sources


def load_gtap_sector_names(gtap_sectors_path: str | Path = DEFAULT_GTAP_SECTORS_PATH) -> dict[str, str]:
    gtap_sectors_path = as_path(gtap_sectors_path)
    if not gtap_sectors_path.exists():
        return {}

    wb = load_workbook(gtap_sectors_path, data_only=True, read_only=True)
    sheet_name = "Finalized GTAP sectors"
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]

    names_by_code: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = row[0] if len(row) > 0 else None
        gtap_entice_name = row[3] if len(row) > 3 else None
        gtap_parent_code = row[4] if len(row) > 4 else None
        gtap_power_name = row[5] if len(row) > 5 else None
        if code is None:
            continue
        name = gtap_entice_name or gtap_power_name
        if name is None:
            continue
        names_by_code[str(code).upper()] = str(name)
        if gtap_parent_code and gtap_power_name:
            names_by_code[str(gtap_parent_code).upper()] = str(gtap_power_name)

    wb.close()
    return names_by_code


def read_sector_templates(source_dir: str | Path) -> list[SectorTemplate]:
    templates: list[SectorTemplate] = []

    for path in workbook_paths(source_dir):
        wb = load_workbook(path, data_only=True, read_only=True)
        summary = wb["Summary"]
        master = wb["Master"]

        summary_sources: list[tuple[str | None, str | None]] = []
        row = 8
        while row <= summary.max_row:
            label = summary.cell(row=row, column=1).value
            value = summary.cell(row=row, column=2).value
            if label is None and value is None:
                break
            summary_sources.append((label, value))
            row += 1

        first_master_row = next(master.iter_rows(min_row=2, max_row=2, values_only=True))
        templates.append(
            SectorTemplate(
                source_path=path,
                sector_code=summary["B2"].value,
                sector_name=sector_name_from_source_path(path),
                parent_sector=summary["B3"].value,
                parent_sector_name=None,
                inventory_version=summary["B4"].value,
                data_collection_lead=summary["B5"].value,
                add_or_split=first_master_row[11] or "Split",
                summary_sources=normalize_summary_sources(summary_sources),
            )
        )
        wb.close()

    return templates


def read_total_outputs(source_path: str | Path, sector_code: str) -> OrderedDict[str, float]:
    wb = load_workbook(as_path(source_path), data_only=True, read_only=True)
    ws = wb["Total outputs"]

    total_outputs: OrderedDict[str, float] = OrderedDict()
    for row_sector, region, quantity, *_ in ws.iter_rows(min_row=2, values_only=True):
        if region is None:
            continue
        if row_sector not in (None, sector_code):
            continue
        total_outputs[region] = float(quantity or 0.0)

    wb.close()
    return total_outputs


def read_purdue_trade_tables(
    splitargs_path: str | Path,
) -> tuple[dict[str, list[tuple[str, str, float]]], str]:
    wb = load_workbook(as_path(splitargs_path), data_only=True, read_only=True)
    ws = wb["BLTTRD"]

    trades_by_sector: dict[str, list[tuple[str, str, float]]] = defaultdict(list)
    for region_from, region_to, sector_code, quantity in ws.iter_rows(min_row=2, values_only=True):
        if region_from is None or region_to is None or sector_code is None or quantity is None:
            continue
        trades_by_sector[str(sector_code).upper()].append(
            (str(region_from).lower(), str(region_to).lower(), float(quantity))
        )

    wb.close()
    return dict(trades_by_sector), as_path(splitargs_path).name


def resolve_export_region_maps(
    db_regions: list[str],
) -> tuple[list[str], OrderedDict[str, list[str]], OrderedDict[str, list[str]]]:
    cluster_members = OrderedDict((region, [region]) for region in db_regions)
    export_members = OrderedDict((region, [region]) for region in db_regions)
    return list(db_regions), cluster_members, export_members


def inventory_sheet_names(workbook) -> list[str]:
    return [name for name in workbook.sheetnames if name not in STATIC_SHEETS]


def clear_sheet(ws) -> None:
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None


def rewrite_summary(
    ws,
    template: SectorTemplate,
    link_sheet_names: list[str],
) -> None:
    ws["A1"] = "INFO"
    ws["A2"] = "Sector"
    ws["B2"] = template.sector_code
    ws["C2"] = template.sector_name
    ws["A3"] = "Parent sector"
    ws["B3"] = template.parent_sector
    ws["C3"] = template.parent_sector_name or template.parent_sector
    ws["A4"] = "Inventory version"
    ws["B4"] = template.inventory_version
    ws["C4"] = None
    ws["A5"] = None
    ws["B5"] = None
    ws["C5"] = None
    ws["A6"] = None
    ws["B6"] = None
    ws["C6"] = None
    ws["A7"] = "Sources"

    row = 8
    for label, value in template.summary_sources:
        ws.cell(row=row, column=1).value = label
        ws.cell(row=row, column=2).value = value
        ws.cell(row=row, column=3).value = None
        row += 1

    while row <= ws.max_row:
        for col_idx in range(1, 4):
            ws.cell(row=row, column=col_idx).value = None
        row += 1

    for row_idx in range(1, max(ws.max_row, len(link_sheet_names)) + 5):
        ws.cell(row=row_idx, column=9).value = None

    ws.cell(row=1, column=9, value="LINKS TO PAGES")
    for idx, sheet_name in enumerate(link_sheet_names, start=2):
        ws.cell(row=idx, column=9).value = sheet_name


def rewrite_master(
    ws,
    template: SectorTemplate,
    aggregated_regions: list[str],
) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for region in aggregated_regions:
        sheet_name = f"{template.sector_code}_{region}"
        ws.append(
            [
                region,
                template.sector_code,
                sheet_name,
                1,
                "M USD",
                None,
                None,
                template.parent_sector,
                None,
                None,
                None,
                template.add_or_split,
            ]
        )


def rewrite_region_clusters(
    ws,
    global_regions: list[str],
    cluster_members: OrderedDict[str, list[str]],
) -> None:
    clear_sheet(ws)
    columns: list[tuple[str, list[str]]] = [("GLOBAL", list(global_regions))]
    columns.extend(
        (cluster_name, members)
        for cluster_name, members in cluster_members.items()
        if len(members) > 1
    )

    for col_idx, (cluster_name, members) in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=cluster_name)
        for row_idx, member in enumerate(members, start=2):
            ws.cell(row=row_idx, column=col_idx, value=member)


def rewrite_sectors_clusters(ws) -> None:
    clear_sheet(ws)


def rewrite_db_units(ws, units_dict) -> None:
    clear_sheet(ws)
    ws.cell(row=1, column=3, value="unit")
    row_idx = 2

    for item_type in ("Sector", "Factor of production", "Satellite account"):
        if item_type not in units_dict:
            continue
        frame = units_dict[item_type]
        for item, unit in frame["unit"].items():
            ws.cell(row=row_idx, column=1, value=item_type)
            ws.cell(row=row_idx, column=2, value=item)
            ws.cell(row=row_idx, column=3, value=unit)
            row_idx += 1


def rewrite_trades(
    ws,
    template: SectorTemplate,
    trade_rows: list[list[object]],
) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    headers = ["Sector_from", "Region_from", "Region_to", "Quantity", "Unit", "Source", "Notes"]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, row in enumerate(trade_rows, start=2):
        ws.cell(row=row_idx, column=1, value=template.sector_code)
        ws.cell(row=row_idx, column=2, value=row[0])
        ws.cell(row=row_idx, column=3, value=row[1])
        ws.cell(row=row_idx, column=4, value=row[2])
        ws.cell(row=row_idx, column=5, value=row[3])
        ws.cell(row=row_idx, column=6, value=row[4])
        ws.cell(row=row_idx, column=7, value=row[5])


def rewrite_total_outputs(
    ws,
    template: SectorTemplate,
    total_outputs: OrderedDict[str, float],
    sector_unit: str,
) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    headers = ["Sector", "Region", "Quantity", "Unit", "Source", "Notes"]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, (region, quantity) in enumerate(total_outputs.items(), start=2):
        ws.cell(row=row_idx, column=1, value=template.sector_code)
        ws.cell(row=row_idx, column=2, value=region)
        ws.cell(row=row_idx, column=3, value=float(quantity))
        ws.cell(row=row_idx, column=4, value=sector_unit)


def reset_inventory_sheet(ws) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def is_effectively_zero(value: float, tolerance: float) -> bool:
    return math.isnan(value) or abs(value) <= tolerance


def build_inventory_rows(
    z_coeffs,
    v_coeffs,
    e_coeffs,
    units_lookup: dict[str, dict[str, str]],
    sector_order: list[str],
    factor_order: list[str],
    satellite_order: list[str],
    tolerance: float,
) -> list[list[object]]:
    rows: list[list[object]] = []

    def append_rows(order: list[str], values, item_type: str) -> None:
        units = units_lookup[item_type]
        for item in order:
            value = values.get(item)
            if value is None or is_effectively_zero(float(value), tolerance):
                continue
            rows.append([float(value), units[item], item, item_type, item, "GLOBAL", "Update"])

    append_rows(sector_order, z_coeffs, "Sector")
    append_rows(factor_order, v_coeffs, "Factor of production")
    append_rows(satellite_order, e_coeffs, "Satellite account")
    return rows


def render_inventory_rows(
    inventory_payload: OrderedDict[str, dict[str, object]],
    units_lookup: dict[str, dict[str, str]],
    sector_order: list[str],
    factor_order: list[str],
    satellite_order: list[str],
    tolerance: float,
) -> tuple[list[tuple[str, list[list[object]]]], list[str]]:
    rendered_inventories: list[tuple[str, list[list[object]]]] = []
    empty_regions: list[str] = []

    for region, payload in inventory_payload.items():
        rows = build_inventory_rows(
            payload["z"],
            payload["v"],
            payload["e"],
            units_lookup,
            sector_order,
            factor_order,
            satellite_order,
            tolerance,
        )
        if rows:
            rendered_inventories.append((region, rows))
        else:
            empty_regions.append(region)

    return rendered_inventories, empty_regions


def build_trade_rows(
    sector_code: str,
    trade_payload: list[tuple[str, str, float]] | None,
    exported_regions: Iterable[str],
    tolerance: float,
) -> list[list[object]]:
    if not trade_payload:
        return []

    exported_region_codes = {region.lower() for region in exported_regions}
    rows: list[list[object]] = []
    for region_from, region_to, quantity in trade_payload:
        if region_from == region_to:
            continue
        if is_effectively_zero(float(quantity), tolerance):
            continue
        if region_from not in exported_region_codes or region_to not in exported_region_codes:
            continue
        rows.append([region_from, region_to, float(quantity), "M USD", "COMTRADE and BACI", None])

    return rows


def collect_non_sum_to_1_inventories(
    rendered_inventories: list[tuple[str, list[list[object]]]],
    sum_check_tolerance: float,
) -> list[str]:
    issues: list[str] = []

    for region, rows in rendered_inventories:
        total = sum(float(row[0]) for row in rows)
        if abs(total - 1.0) <= sum_check_tolerance:
            continue
        issues.append(f"{region} ({total:.12g})")

    return issues


def write_export_report(
    output_dir: Path,
    splitargs_path: Path,
    missing_trade_sectors: list[str],
    null_inventories: OrderedDict[str, list[str]],
    non_sum_to_1_inventories: OrderedDict[str, list[str]],
) -> Path:
    report_path = output_dir / REPORT_FILENAME

    lines = [
        "D2.4 export report",
        "",
        f"Purdue trade source: {splitargs_path}",
        "",
        "Missing trade data by sector:",
    ]
    if missing_trade_sectors:
        lines.extend(f"- {sector_code}" for sector_code in missing_trade_sectors)
    else:
        lines.append("- none")

    lines.extend([
        "",
        "Null inventories by sector:",
    ])
    for sector_code, regions in null_inventories.items():
        if regions:
            lines.append(f"- {sector_code}: {', '.join(regions)}")
        else:
            lines.append(f"- {sector_code}: none")

    lines.extend([
        "",
        "Non-sum-to-1 inventories by sector:",
    ])
    for sector_code, regions in non_sum_to_1_inventories.items():
        if regions:
            lines.append(f"- {sector_code}: {', '.join(regions)}")
        else:
            lines.append(f"- {sector_code}: none")

    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return report_path


def aggregate_sector_coefficients(
    db,
    sector_code: str,
    regions: Iterable[str],
):
    z_coefficients = db.z
    v_coefficients = db.v
    e_coefficients = db.e

    sector_rows = z_coefficients.index.get_level_values("Level") == "Sector"
    sector_block = z_coefficients.loc[sector_rows]
    sector_columns = z_coefficients.columns

    results: OrderedDict[str, dict[str, object]] = OrderedDict()

    for region in regions:
        column = (region, "Sector", sector_code)

        if column not in sector_columns:
            results[region] = {
                "z": OrderedDict(),
                "v": OrderedDict(),
                "e": OrderedDict(),
            }
            continue

        z_coeffs = sector_block.loc[:, column].groupby(level="Item").sum()
        v_coeffs = v_coefficients.loc[:, column]
        e_coeffs = e_coefficients.loc[:, column]

        results[region] = {
            "z": z_coeffs.to_dict(),
            "v": v_coeffs.to_dict(),
            "e": e_coeffs.to_dict(),
        }

    return results


def reorder_sheets(workbook, ordered_names: list[str]) -> None:
    workbook._sheets = [workbook[name] for name in ordered_names]


def create_output_workbook(
    template: SectorTemplate,
    output_dir: Path,
    global_regions: list[str],
    cluster_members: OrderedDict[str, list[str]],
    units_dict,
    total_outputs: OrderedDict[str, float],
    rendered_inventories: list[tuple[str, list[list[object]]]],
    trade_rows: list[list[object]],
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{template.sector_code} - {template.sector_name}.xlsx"

    wb = load_workbook(template.source_path)
    inventory_templates = inventory_sheet_names(wb)
    if not inventory_templates:
        raise RuntimeError(f"No inventory sheet found in template workbook '{template.source_path}'.")

    template_sheet_name = inventory_templates[0]
    for sheet_name in inventory_templates[1:]:
        del wb[sheet_name]

    inventory_template = wb[template_sheet_name]
    inventory_template.title = "_INVENTORY_TEMPLATE"

    summary_ws = wb["Summary"]
    master_ws = wb["Master"]
    region_clusters_ws = wb["Regions Clusters"]
    sectors_clusters_ws = wb["Sectors Clusters"]
    trades_ws = wb["Trades"] if "Trades" in wb.sheetnames else wb.create_sheet("Trades")
    db_units_ws = wb["DB units"]
    total_outputs_ws = wb["Total outputs"]

    rewrite_region_clusters(region_clusters_ws, global_regions, cluster_members)
    rewrite_sectors_clusters(sectors_clusters_ws)
    rewrite_db_units(db_units_ws, units_dict)
    aggregated_regions = [region for region, _ in rendered_inventories]
    desired_inventory_sheet_names = [f"{template.sector_code}_{region}" for region in aggregated_regions]

    rewrite_summary(
        summary_ws,
        template,
        [
            "Master",
            "Regions Clusters",
            "Sectors Clusters",
            "Factors Clusters",
            "Total outputs",
            "Trades",
            *desired_inventory_sheet_names,
            "DB units",
        ],
    )
    rewrite_master(master_ws, template, aggregated_regions)
    rewrite_trades(trades_ws, template, trade_rows)
    rewrite_total_outputs(
        total_outputs_ws,
        template,
        OrderedDict((region, total_outputs[region]) for region in aggregated_regions),
        units_dict["Sector"].loc[template.sector_code, "unit"],
    )

    built_inventory_names: list[str] = []
    for idx, (region, rows) in enumerate(rendered_inventories):
        sheet_name = f"{template.sector_code}_{region}"
        if idx == 0:
            ws = inventory_template
            ws.title = sheet_name
        else:
            ws = wb.copy_worksheet(inventory_template)
            ws.title = sheet_name

        reset_inventory_sheet(ws)
        for row in rows:
            ws.append(row)
        built_inventory_names.append(sheet_name)

    if "_INVENTORY_TEMPLATE" in wb.sheetnames:
        del wb["_INVENTORY_TEMPLATE"]

    ordered_names = [
        "Summary",
        "Master",
        "Regions Clusters",
        "Sectors Clusters",
        "Factors Clusters",
        "Total outputs",
        "Trades",
        *built_inventory_names,
        "DB units",
    ]
    reorder_sheets(wb, ordered_names)
    wb.save(output_path)
    wb.close()
    return output_path


def export_d24_inventories(
    db,
    source_dir: str | Path = DEFAULT_SOURCE_DIR,
    output_dir: str | Path = DEFAULT_OUTPUT_DIR,
    purdue_splitargs_path: str | Path = DEFAULT_PURDUE_SPLITARGS_PATH,
    tolerance: float = 1e-12,
    inventory_sum_check_tolerance: float = DEFAULT_INVENTORY_SUM_CHECK_TOLERANCE,
) -> list[Path]:
    source_dir = as_path(source_dir)
    output_dir = as_path(output_dir)
    purdue_splitargs_path = as_path(purdue_splitargs_path)

    templates = read_sector_templates(source_dir)
    if not templates:
        raise RuntimeError(f"No .xlsx sector inventory workbooks found in '{source_dir}'.")

    db_regions = list(db.get_index("Region"))
    global_regions, cluster_members, export_members = resolve_export_region_maps(db_regions)

    units_dict = db.units
    sector_order = units_dict["Sector"].index.tolist()
    factor_order = (
        units_dict["Factor of production"].index.tolist()
        if "Factor of production" in units_dict
        else []
    )
    satellite_order = (
        units_dict["Satellite account"].index.tolist()
        if "Satellite account" in units_dict
        else []
    )
    units_lookup = {
        item_type: frame["unit"].to_dict()
        for item_type, frame in units_dict.items()
    }
    trades_by_sector, trade_source_name = read_purdue_trade_tables(purdue_splitargs_path)
    sector_names = load_gtap_sector_names()
    for template in templates:
        sector_names[template.sector_code] = template.sector_name
    for template in templates:
        template.parent_sector_name = sector_names.get(template.parent_sector, template.parent_sector)

    written_files: list[Path] = []
    missing_trade_sectors: list[str] = []
    null_inventories: OrderedDict[str, list[str]] = OrderedDict()
    non_sum_to_1_inventories: OrderedDict[str, list[str]] = OrderedDict()
    for template in templates:
        total_outputs = read_total_outputs(template.source_path, template.sector_code)
        inventory_payload = aggregate_sector_coefficients(db, template.sector_code, total_outputs.keys())
        rendered_inventories, empty_regions = render_inventory_rows(
            inventory_payload,
            units_lookup,
            sector_order,
            factor_order,
            satellite_order,
            tolerance,
        )
        null_inventories[template.sector_code] = empty_regions
        non_sum_to_1_inventories[template.sector_code] = collect_non_sum_to_1_inventories(
            rendered_inventories,
            inventory_sum_check_tolerance,
        )
        if template.sector_code not in trades_by_sector:
            missing_trade_sectors.append(template.sector_code)
        trade_rows = build_trade_rows(
            template.sector_code,
            trades_by_sector.get(template.sector_code),
            [region for region, _ in rendered_inventories],
            tolerance,
        )
        output_path = create_output_workbook(
            template=template,
            output_dir=output_dir,
            global_regions=global_regions,
            cluster_members=cluster_members,
            units_dict=units_dict,
            total_outputs=total_outputs,
            rendered_inventories=rendered_inventories,
            trade_rows=trade_rows,
        )
        written_files.append(output_path)

    write_export_report(
        output_dir,
        purdue_splitargs_path,
        missing_trade_sectors,
        null_inventories,
        non_sum_to_1_inventories,
    )
    return written_files


def main() -> None:
    args = parse_args()
    warnings.filterwarnings("ignore")

    mario = ensure_mario_import(args.mario_src)

    print(f"Loading GTAP database from {args.db_path}")
    db = mario.parse_from_parquet(str(args.db_path), table="IOT", mode="flows")

    print(f"Reading add-sector workbooks from {args.source_dir}")
    db.read_add_sectors_excel(path=str(args.source_dir), read_inventories=True, split=False)

    print("Extending database with MARIO sectors")
    db.add_sectors(split=False, VA_fix=True)

    print(f"Writing D2.4 inventories to {args.output_dir}")
    written_files = export_d24_inventories(
        db,
        source_dir=args.source_dir,
        output_dir=args.output_dir,
        purdue_splitargs_path=args.purdue_splitargs_path,
        tolerance=args.tolerance,
        inventory_sum_check_tolerance=args.inventory_sum_check_tolerance,
    )
    for output_path in written_files:
        print(f"  wrote {output_path.name}")

    print(f"Completed: {len(written_files)} workbook(s) created in {args.output_dir}")


if __name__ == "__main__":
    main()
