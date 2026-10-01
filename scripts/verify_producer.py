"""Exercise coefficient insertion, workbook export and an actual CVXLAB split.

Requires the isolated producer environment. With --baseline-dir, retain one GTAP
region and aggregate all others to ROW, without aggregating sectors. All generated
data stay in a new output directory. This is a representative integration check,
not a reconstruction of the complete historical 163-region run.
"""
from __future__ import annotations

import argparse
from collections import OrderedDict
import hashlib
import json
import logging
import math
from pathlib import Path
import platform
import sys
import traceback
import warnings

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import mario
import numpy as np
import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook
from mario.ops.add_sector_specs import ADD_SECTOR_SPLIT_EXCLUSION_COLUMNS

from entice_inventory.export.build_d24_inventories import (
    aggregate_sector_coefficients,
    create_output_workbook,
    read_sector_templates,
    render_inventory_rows,
    save_workbook_atomic,
)


def digest(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def grouped_index(index, region):
    if not isinstance(index, pd.MultiIndex) or "Region" not in index.names:
        return index
    level = index.names.index("Region")
    labels = []
    for label in index:
        parts = list(label)
        parts[level] = region if parts[level] == region else "ROW"
        labels.append(tuple(parts))
    return pd.MultiIndex.from_tuples(labels, names=index.names)


def aggregate_parquet(path, region, chunk_size=256):
    """Aggregate in column chunks; do not materialise the full global Z matrix."""
    parquet = pq.ParquetFile(path)
    metadata = json.loads(parquet.schema_arrow.metadata[b"pandas"])
    indexes = [c for c in metadata["index_columns"] if isinstance(c, str)]
    columns = [c["field_name"] for c in metadata["columns"] if c["field_name"] not in indexes]
    pieces = []
    for offset in range(0, len(columns), chunk_size):
        block = parquet.read(columns=columns[offset:offset + chunk_size] + indexes).to_pandas()
        block.index = grouped_index(block.index, region)
        if isinstance(block.index, pd.MultiIndex):
            block = block.groupby(level=list(range(block.index.nlevels)), sort=False).sum()
        block.columns = grouped_index(block.columns, region)
        if isinstance(block.columns, pd.MultiIndex):
            block = block.T.groupby(level=list(range(block.columns.nlevels)), sort=False).sum().T
        pieces.append(block)
    result = pd.concat(pieces, axis=1)
    if isinstance(result.columns, pd.MultiIndex):
        result = result.T.groupby(level=list(range(result.columns.nlevels)), sort=False).sum().T
    return result.sort_index().sort_index(axis=1)


def read_baseline(path, region):
    matrices = {}
    for name in ("Z", "Y", "V", "E", "EY", "VY"):
        source = path / f"{name}.parquet"
        if source.exists():
            print(f"Aggregating {name}", flush=True)
            matrices[name] = aggregate_parquet(source, region)
    frame = pd.read_parquet(path / "units.parquet")
    units = {kind: part.droplevel(0) for kind, part in frame.groupby(level=0, sort=False)}
    return mario.Database(name=f"GTAP pilot: {region} and ROW", table="IOT", units=units, **matrices)


def metric(observed, expected, absolute, relative):
    """Check elementwise tolerances; include non-finite values as failures."""
    if isinstance(observed, (pd.Series, pd.DataFrame)):
        observed, expected = observed.align(expected, join="outer")
    actual = np.asarray(observed, dtype=float)
    target = np.asarray(expected, dtype=float)
    finite = np.isfinite(actual) & np.isfinite(target)
    residual = np.abs(actual - target)
    allowance = absolute + relative * np.abs(target)
    failures = ~finite | (residual > allowance)
    examples = []
    for location in np.argwhere(failures)[:5]:
        key = tuple(location)
        coordinates = {}
        if isinstance(observed, (pd.Series, pd.DataFrame)):
            coordinates["row"] = str(observed.index[location[0]])
        if isinstance(observed, pd.DataFrame):
            coordinates["column"] = str(observed.columns[location[1]])
        examples.append({
            **coordinates,
            "observed": float(actual[key]) if np.isfinite(actual[key]) else None,
            "expected": float(target[key]) if np.isfinite(target[key]) else None,
            "absolute_residual": float(residual[key]) if finite[key] else None,
            "allowance": float(allowance[key]) if np.isfinite(allowance[key]) else None,
        })
    return {
        "passed": not bool(np.any(failures)),
        "failed_entries": int(np.count_nonzero(failures)),
        "non_finite_entries": int(np.count_nonzero(~finite)),
        "maximum_absolute_residual": float(np.max(residual[finite])) if finite.any() else None,
        "absolute_tolerance": absolute,
        "relative_tolerance": relative,
        "failure_examples": examples,
    }


def accounting(db, absolute, relative):
    x = db.X.iloc[:, 0]
    return {
        "row_accounting": metric(db.Z.sum(axis=1) + db.Y.sum(axis=1), x, absolute, relative),
        "column_accounting": metric(db.Z.sum(axis=0) + db.V.sum(axis=0), x, absolute, relative),
        "minimum_Z": float(db.Z.min().min()),
        "minimum_Y": float(db.Y.min().min()),
        "minimum_V": float(db.V.min().min()),
    }


def trim_source(source, target, sector, region):
    """Select an existing regional profile; retain its cost/cluster definitions."""
    workbook = load_workbook(source)
    master = workbook["Master"]
    rows = list(master.values)
    header = list(rows[0])
    selected = [r for r in rows[1:] if r[header.index("Region")] == region and r[header.index("Sector")] == sector]
    if len(selected) != 1:
        raise ValueError(f"Expected one {sector}/{region} Master row, found {len(selected)}")
    sheet = selected[0][header.index("Inventory sheet")]
    inventory_names = {r[header.index("Inventory sheet")] for r in rows[1:] if r[header.index("Inventory sheet")]}
    for name in inventory_names - {sheet}:
        if name in workbook:
            del workbook[name]
    master.delete_rows(1, master.max_row)
    master.append(header)
    master.append(selected[0])
    outputs = workbook["Total outputs"]
    output_rows = list(outputs.values)
    outputs.delete_rows(1, outputs.max_row)
    outputs.append(output_rows[0])
    for row in output_rows[1:]:
        if row[0] == sector and row[1] == region:
            outputs.append(row)
    if outputs.max_row != 2:
        raise ValueError("Exactly one output observation is required for the pilot")
    # The pilot maps group members to its aggregated geography. A partially
    # covered ROW becomes the whole ROW here; the native reference separately
    # detects omitted origins. This is a diagnostic aggregation, not a migration
    # rule for production inventories.
    clusters = workbook["Regions Clusters"]
    old = list(clusters.values)
    mapped = {}
    for column, name in enumerate(old[0]):
        if name:
            mapped[name] = list(dict.fromkeys(region if r[column] == region else "ROW" for r in old[1:] if column < len(r) and r[column]))
    clusters.delete_rows(1, clusters.max_row)
    for column, (name, members) in enumerate(mapped.items(), 1):
        clusters.cell(1, column, name)
        for row, member in enumerate(members, 2):
            clusters.cell(row, column, member)
    if "Trades" in workbook:
        trades = workbook["Trades"]
        if trades.max_row > 1:
            trades.delete_rows(2, trades.max_row - 1)
    save_workbook_atomic(workbook, target)
    workbook.close()
    return float(output_rows[[r[1] for r in output_rows].index(region)][2])


def export_sector(db, source, sector, region, output_dir):
    templates, skipped = read_sector_templates(source.parent, {}, [sector], [source])
    if len(templates) != 1 or skipped:
        raise ValueError(f"Cannot select template: {skipped}")
    template = templates[0]
    units = db.units
    payload = aggregate_sector_coefficients(db, sector, [region])
    rows, empty = render_inventory_rows(
        payload, {kind: frame["unit"].to_dict() for kind, frame in units.items()},
        *[units[k].index.tolist() for k in ("Sector", "Factor of production", "Satellite account")], 1e-12,
    )
    if empty:
        raise ValueError(f"Empty inventories: {empty}")
    workbook = load_workbook(source, read_only=True, data_only=True)
    outputs = OrderedDict((r[1], float(r[2])) for r in workbook["Total outputs"].iter_rows(min_row=2, values_only=True) if r[0] == sector)
    workbook.close()
    return create_output_workbook(template, output_dir, db.get_index("Region"), OrderedDict(), units, outputs, rows, [])


def compare_published(exported, published, sheet):
    def records(path):
        workbook = load_workbook(path, read_only=True, data_only=True)
        result = {}
        for row in workbook[sheet].iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            key = tuple(row[1:7])
            result[key] = result.get(key, 0.0) + float(row[0])
        workbook.close()
        return result
    actual, baseline = records(exported), records(published)
    differences = []
    for key in sorted(actual.keys() | baseline.keys(), key=str):
        a, b = actual.get(key, 0.0), baseline.get(key, 0.0)
        if not math.isclose(a, b, rel_tol=1e-10, abs_tol=1e-12):
            differences.append({"key": key, "current": a, "published": b, "difference": a - b})
    return {"current_records": len(actual), "published_records": len(baseline), "differences": differences}


def native_hye_reference(baseline, source, published, region):
    """Independently replay this HYE profile's formulas at native geography.

    This narrowly scoped cross-check explains differences caused by aggregating
    before insertion. It does not replace the MARIO producer or reconstruct other
    sectors. In particular, uncovered origins retain inherited parent inputs.
    """
    x = pd.read_parquet(baseline / "X.parquet").iloc[:, 0]
    native_regions = set(x.index.get_level_values("Region"))
    if len(native_regions) <= 2:
        return {"status": "skipped", "reason": "Native-geography reference requires the original GTAP baseline."}
    workbook = load_workbook(source, read_only=True, data_only=True)
    def clusters(name):
        rows = list(workbook[name].values)
        return {label: [row[c] for row in rows[1:] if c < len(row) and row[c]] for c, label in enumerate(rows[0]) if label}
    regions, sectors, factors = [clusters(name) for name in ("Regions Clusters", "Sectors Clusters", "Factors Clusters")]
    inventory = list(workbook[f"HYE_{region}"].values)[1:]
    parent = workbook["Summary"]["B3"].value
    workbook.close()
    column = (region, "Sector", parent)
    z = pq.ParquetFile(baseline / "Z.parquet").read(columns=[str(column), "Region", "Level", "Item"]).to_pandas().iloc[:, 0]
    v = pd.read_parquet(baseline / "V.parquet").div(x.replace(0, np.nan), axis=1).fillna(0)
    e = pq.ParquetFile(baseline / "E.parquet").read(columns=[str(column), "Item"]).to_pandas().iloc[:, 0] / float(x.loc[column])
    missing = sorted(native_regions - set(regions["GLOBAL"]))
    inherited = z.loc[pd.IndexSlice[missing, "Sector", :]].groupby(level="Item").sum() / float(x.loc[column]) if missing else pd.Series(dtype=float)
    expected = {("Sector", item): float(inherited.get(item, 0)) for item in x.index.get_level_values("Item").unique()}
    expected.update({("Factor of production", item): 0.0 for item in v.index})
    expected.update({("Satellite account", item): float(value) for item, value in e.items()})
    for row in inventory:
        if not row or row[0] is None:
            continue
        quantity, _, _, kind, item, origin, change = row[:7]
        if origin != "GLOBAL" or change != "Update":
            raise ValueError("The independent HYE reference only handles GLOBAL Update rows")
        if kind == "Sector":
            members = sectors.get(item, [item])
            selected = z.loc[pd.IndexSlice[regions[origin], "Sector", members]]
            weights = selected.groupby(level="Item").sum() / selected.sum()
        elif kind == "Factor of production":
            weights = v.loc[factors.get(item, [item])].sum(axis=1)
            weights /= weights.sum()
        else:
            raise ValueError(f"Unsupported HYE reference row type: {kind}")
        if not np.isfinite(weights).all():
            raise ValueError(f"Undefined HYE reference weights for {item}")
        for item, weight in weights.items():
            expected[(kind, item)] += float(quantity) * float(weight)
    workbook = load_workbook(published, read_only=True, data_only=True)
    actual = {(r[3], r[4]): float(r[0]) for r in workbook[f"HYE_{region}"].iter_rows(min_row=2, values_only=True) if r[0] is not None}
    workbook.close()
    differences = []
    for key in sorted(actual.keys() | expected.keys()):
        observed, predicted = actual.get(key, 0), expected.get(key, 0)
        if not math.isclose(observed, predicted, abs_tol=1e-12, rel_tol=1e-10):
            differences.append({"key": key, "published": observed, "replayed": predicted})
    return {
        "status": "passed" if not differences else "differences_found",
        "scope": f"Independent formula replay for HYE/{region}; native geography, not a complete database build.",
        "published_records": len(actual), "differences": differences,
        "maximum_absolute_difference": max(abs(actual.get(k, 0) - expected.get(k, 0)) for k in actual.keys() | expected.keys()),
        "GLOBAL_missing_regions": missing,
        "factor_weight_rule": "Unweighted sums of value-added coefficients across all baseline columns; these weights change when geography is aggregated first.",
        "uncovered_origin_rule": "Parent coefficients remain for origins outside the inventory's GLOBAL cluster.",
    }


def scaled_database(db, scale):
    return mario.Database(
        name="Numerically scaled split input", table="IOT", units=db.units,
        Z=db.Z / scale, Y=db.Y / scale, V=db.V / scale, VY=db.VY / scale,
        E=db.E, EY=db.EY,
    )


def split_workbook(source, target, scale, sector):
    workbook = load_workbook(source)
    for sheet in ("Total outputs", "Trades"):
        if sheet not in workbook:
            continue
        ws = workbook[sheet]
        header = [cell.value for cell in ws[1]]
        quantity = header.index("Quantity") + 1
        for row in range(2, ws.max_row + 1):
            if ws.cell(row, quantity).value is not None:
                ws.cell(row, quantity).value /= scale
    for ws in workbook:
        if ws.title.startswith(f"{sector}_"):
            for row in range(2, ws.max_row + 1):
                if ws.cell(row, 4).value == "Satellite account" and ws.cell(row, 1).value is not None:
                    ws.cell(row, 1).value *= scale
    if "Tolerances" in workbook:
        del workbook["Tolerances"]
    ws = workbook.create_sheet("Tolerances")
    ws.append(["tol_Name", "values"])
    ws.append(["delta", 0.0])
    ws.append(["eps", 0.0])
    if "Exclusions" not in workbook:
        workbook.create_sheet("Exclusions").append(list(ADD_SECTOR_SPLIT_EXCLUSION_COLUMNS.values()))
    save_workbook_atomic(workbook, target)
    workbook.close()


def collapsed(frame, child, parent):
    frame = frame.copy()
    for axis in (0, 1):
        index = frame.index if axis == 0 else frame.columns
        if not isinstance(index, pd.MultiIndex) or "Item" not in index.names:
            continue
        labels = [tuple(parent if i == index.names.index("Item") and x == child else x for i, x in enumerate(label)) for label in index]
        replacement = pd.MultiIndex.from_tuples(labels, names=index.names)
        if axis == 0:
            frame.index = replacement
            frame = frame.groupby(level=list(range(index.nlevels))).sum()
        else:
            frame.columns = replacement
            frame = frame.T.groupby(level=list(range(index.nlevels))).sum().T
    return frame


class SolverStatus(logging.Handler):
    def __init__(self):
        super().__init__()
        self.statuses = []

    def emit(self, record):
        message = record.getMessage()
        if "Problem status:" in message:
            self.statuses.append(message.split("Problem status:", 1)[1].strip(" '\""))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline-dir", required=True, type=Path)
    parser.add_argument("--source-workbook", required=True, type=Path)
    parser.add_argument("--published-workbook", type=Path)
    parser.add_argument("--sector", default="HYE", choices=["HYE"])
    parser.add_argument("--region", default="DEU")
    parser.add_argument("--scale", type=float, default=100.0)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--absolute-tolerance", type=float, default=1e-6)
    parser.add_argument("--relative-tolerance", type=float, default=1e-8)
    args = parser.parse_args()
    for name in ("scale", "absolute_tolerance", "relative_tolerance"):
        value = getattr(args, name)
        if not math.isfinite(value) or value < 0 or (name == "scale" and value == 0):
            parser.error(f"Invalid {name}: {value}")
    args.output_dir.mkdir(parents=True, exist_ok=False)
    result = {"status": "running", "scope": "One region plus aggregated ROW; all baseline sectors. No full historical-build certification.",
              "arguments": {k: str(v) if isinstance(v, Path) else v for k, v in vars(args).items()},
              "python": platform.python_version(), "warnings": [], "solver_statuses": []}
    status_handler = SolverStatus()
    logging.getLogger("Model.core.problem").addHandler(status_handler)
    caught = []
    exit_code = 0
    try:
        result["source_sha256"] = digest(args.source_workbook)
        result["baseline_hashes"] = {p.name: digest(p) for p in args.baseline_dir.glob("*.parquet")}
        if args.published_workbook:
            result["published_sha256"] = digest(args.published_workbook)
            result["native_reference"] = native_hye_reference(args.baseline_dir, args.source_workbook, args.published_workbook, args.region)
        result["aggregation_limitations"] = [
            "Factor-cluster weights depend on geography; aggregating before insertion can change coefficients.",
            "Mapping a source cluster that only partly covers ROW expands that cluster to ROW in this diagnostic. Resolve source profiles before aggregation for production use.",
        ]
        db = read_baseline(args.baseline_dir, args.region)
        result["baseline_accounting"] = accounting(db, args.absolute_tolerance, args.relative_tolerance)
        result["baseline_policy"] = "Preserve baseline values; diagnose existing residuals without correcting them."
        db.to_parquet(str(args.output_dir / "baseline"), flows=True, coefficients=False)
        source = args.output_dir / args.source_workbook.name
        target = trim_source(args.source_workbook, source, args.sector, args.region)
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            db.read_add_sectors_excel(source, read_inventories=True, split=False)
            inserted = db.add_sectors(inplace=False, split=False, VA_fix=True, accept_non_unitary_sum=True, ignore_warnings=False)
            inserted.to_parquet(str(args.output_dir / "inserted"), coefficients=True, flows=False)
            exported = export_sector(inserted, source, args.sector, args.region, args.output_dir / "export")
            result["export"] = {"path": str(exported), "sha256": digest(exported)}
            if args.published_workbook:
                result["published_comparison"] = compare_published(exported, args.published_workbook, f"{args.sector}_{args.region}")
            scaled = scaled_database(db, args.scale)
            solver_input = args.output_dir / "split-input.xlsx"
            split_workbook(exported, solver_input, args.scale, args.sector)
            model_dir = args.output_dir / "model"
            model_dir.mkdir()
            scaled.read_add_sectors_excel(solver_input, read_inventories=True, split=True)
            solver_parameters = {"tol_gap_abs": 1e-8, "tol_gap_rel": 1e-8, "tol_feas": 1e-10, "max_iter": 300}
            result["solver"] = "CLARABEL"
            result["solver_parameters"] = solver_parameters
            solved = scaled.add_sectors(inplace=False, split=True, VA_fix=True, accept_non_unitary_sum=True,
                                       ignore_warnings=False, solver="CLARABEL", cvxlab_path=model_dir,
                                       solver_parameters=solver_parameters)
            restored = scaled_database(solved, 1 / args.scale)
            result["warnings"] = list(dict.fromkeys(str(w.message) for w in caught))
        result["solver_statuses"] = status_handler.statuses
        result["accounting"] = accounting(restored, args.absolute_tolerance, args.relative_tolerance)
        result["output_target"] = metric(restored.X.loc[(args.region, "Sector", args.sector)].iloc[0], target, args.absolute_tolerance, args.relative_tolerance)
        templates, _ = read_sector_templates(source.parent, {}, [args.sector], [source])
        parent = templates[0].parent_sector
        result["parent_aggregation"] = {name: metric(collapsed(getattr(restored, name), args.sector, parent), getattr(db, name), args.absolute_tolerance, args.relative_tolerance) for name in ("Z", "Y", "V")}
        result["trade_validation"] = {"status": "skipped", "reason": "No child-specific HYE trade observations are supplied; missing trade is not treated as zero."}
        restored.to_parquet(str(args.output_dir / "solved"), flows=True, coefficients=False)
        numerical = [result["accounting"][key]["passed"] for key in ("row_accounting", "column_accounting")]
        numerical += [result["output_target"]["passed"]] + [r["passed"] for r in result["parent_aggregation"].values()]
        result["status"] = "passed" if all(numerical) and status_handler.statuses and all(s == "optimal" for s in status_handler.statuses) else "completed_with_warnings"
    except Exception as exc:
        result.update(status="failed", error=f"{type(exc).__name__}: {exc}", traceback=traceback.format_exc(), solver_statuses=status_handler.statuses)
        exit_code = 1
    finally:
        result["warnings"] = list(dict.fromkeys(str(w.message) for w in caught))
        logging.getLogger("Model.core.problem").removeHandler(status_handler)
        (args.output_dir / "verification.json").write_text(json.dumps(result, indent=2, allow_nan=False) + "\n")
        print(json.dumps({k: v for k, v in result.items() if k not in ("warnings", "baseline_hashes", "published_comparison", "traceback")}, indent=2))
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
